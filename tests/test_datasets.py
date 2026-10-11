"""FASE 7A — Tests de Client Onboarding + Business Data Workspace.

Cubre: workspace, dataset, upload CSV/XLSX, validación, columnas, mapping,
preview, processing, READY, NO_DATA, ERROR, tenant isolation, roles,
permissions, audit, active dataset, reprocessing, DEMO isolation,
Advisor isolation, dashboard states, lenguaje empresarial.

Ejecutar: .venv/bin/python -m unittest tests.test_datasets -v
"""

from __future__ import annotations

import http.cookiejar
import io
import json
import os
import shutil
import sys
import tempfile
import threading
import time
import unittest
import urllib.request

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tenant import TenantStore, create_company, create_user
import webapp.server as server_mod
from http.server import ThreadingHTTPServer

import datasets.store as ds_store_mod
from datasets import (
    DatasetStore, ensure_demo_dataset, check_dataset_access,
    sanitize_base_name, receive_upload, build_review, suggest_mapping,
    confirm_mapping, build_preview, process_dataset,
    DATA_UPLOAD_MAX_BYTES,
)
from datasets.upload import UploadError
from tenant import audit as audit_mod

PASS = "DatasetsTest123"
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STAGE_DIRS = ["datasets", "uploads", "processed", "profiles", "anomalies",
              "findings", "context", "predictions", "prediction_intelligence",
              "prediction_validation", "business_context"]


def _csv_bytes(n=60, **mods):
    lines = ["factura,producto,cantidad,fecha_venta,precio,cliente"]
    for i in range(1, n + 1):
        qty = mods.get("qty", 10)
        if mods.get("negative") and i == 3:
            qty = -5
        lines.append(
            f"F{i:05d},PROD-{i % 5},{qty},2024-{(i % 12) + 1:02d}-15,25.50,CLI-{i % 8}")
    if mods.get("duplicates"):
        lines.append(lines[5])  # fila duplicada
        lines.append(lines[6])
    if mods.get("cancelled"):
        lines.append("C99999,PROD-1,2,2024-06-10,25.50,CLI-1")
    return ("\n".join(lines) + "\n").encode("utf-8")


def _chunks(data: bytes, size: int = 4096):
    for i in range(0, len(data), size):
        yield data[i:i + size]


def _xlsx_bytes(n=30):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Invoice", "StockCode", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID"])
    for i in range(1, n + 1):
        ws.append([f"F{i:05d}", f"SKU-{i % 5}", i % 9 + 1,
                   f"2024-{(i % 12) + 1:02d}-15", 19.99 + i % 10, f"C{i % 8}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cleanup_company_data(company_id: str):
    for stage in STAGE_DIRS:
        path = os.path.join(PROJECT_ROOT, "data", stage, company_id)
        shutil.rmtree(path, ignore_errors=True)


class DatasetStoreTest(unittest.TestCase):
    """Dataset CRUD y estados (unit)."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="zayvero7a_unit_")
        self.orig_uploads = ds_store_mod.UPLOADS_DIR
        ds_store_mod.UPLOADS_DIR = os.path.join(self.tmp, "uploads")

    def tearDown(self):
        ds_store_mod.UPLOADS_DIR = self.orig_uploads
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _store(self, cid="CMP-UNIT-A"):
        return DatasetStore(cid, base_dir=os.path.join(self.tmp, "datasets"))

    def test_01_create_dataset_campos(self):
        st = self._store()
        ds = st.create_dataset(nombre="Ventas 2024", source="csv")
        self.assertTrue(ds.dataset_id.startswith("DS-"))
        self.assertEqual(ds.status, "DATA_UPLOADED")
        self.assertFalse(ds.is_active)
        self.assertEqual(ds.nombre, "Ventas 2024")
        self.assertEqual(ds.company_id, "CMP-UNIT-A")

    def test_02_set_active_requiere_ready(self):
        st = self._store()
        ds = st.create_dataset(nombre="X", source="csv")
        with self.assertRaises(ValueError):
            st.set_active(ds.dataset_id)

    def test_03_set_active_solo_uno(self):
        st = self._store()
        a = st.create_dataset(nombre="A", source="csv")
        b = st.create_dataset(nombre="B", source="csv")
        for d in (a, b):
            rec = st.get_dataset(d.dataset_id)
            rec.status = "READY"
            st.save_dataset(rec)
        st.set_active(a.dataset_id)
        self.assertEqual(st.get_active().dataset_id, a.dataset_id)
        st.set_active(b.dataset_id)
        self.assertEqual(st.get_active().dataset_id, b.dataset_id)
        self.assertFalse(st.get_dataset(a.dataset_id).is_active)

    def test_04_dataset_state_no_data(self):
        st = self._store(cid="CMP-UNIT-VACIA")
        self.assertEqual(st.dataset_state(), "NO_DATA")

    def test_05_sanitize_base_name(self):
        self.assertEqual(sanitize_base_name("Ventas 2024 Q1.xlsx"), "ventas-2024-q1")
        self.assertTrue(len(sanitize_base_name("ÁÉÍ.csv")) > 0)

    def test_06_check_dataset_access_ajeno(self):
        st = self._store()
        ds = st.create_dataset(nombre="X", source="csv")
        with self.assertRaises(PermissionError):
            check_dataset_access("OTRA-EMPRESA", ds)
        with self.assertRaises(PermissionError):
            check_dataset_access("OTRA-EMPRESA", None)

    def test_07_limite_100mb(self):
        self.assertEqual(DATA_UPLOAD_MAX_BYTES, 100 * 1024 * 1024)


class UploadTest(unittest.TestCase):
    """Carga de archivos (unit)."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="zayvero7a_up_")
        self.orig_uploads = ds_store_mod.UPLOADS_DIR
        ds_store_mod.UPLOADS_DIR = os.path.join(self.tmp, "uploads")

    def tearDown(self):
        ds_store_mod.UPLOADS_DIR = self.orig_uploads
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _store(self):
        return DatasetStore("CMP-UP", base_dir=os.path.join(self.tmp, "datasets"))

    def test_08_upload_csv_ok(self):
        data = _csv_bytes()
        ds = receive_upload(self._store(), filename="ventas.csv",
                            size_bytes=len(data), content_iter=_chunks(data),
                            user_id="USR-1", nombre="Ventas")
        self.assertEqual(ds.status, "DATA_UPLOADED")
        self.assertTrue(os.path.isfile(ds.upload["path"]))
        # Archivo intacto: el contenido original se conserva.
        with open(ds.upload["path"], "rb") as fh:
            self.assertEqual(fh.read(), data)

    def test_09_upload_xlsx_ok(self):
        data = _xlsx_bytes()
        ds = receive_upload(self._store(), filename="reporte.xlsx",
                            size_bytes=len(data), content_iter=_chunks(data),
                            user_id="USR-1")
        self.assertEqual(ds.upload["format"], "xlsx")

    def test_10_upload_extension_invalida(self):
        with self.assertRaises(UploadError) as cm:
            receive_upload(self._store(), filename="reporte.pdf",
                           size_bytes=100, content_iter=_chunks(b"x"),
                           user_id="USR-1")
        self.assertIn("CSV", str(cm.exception))

    def test_11_upload_sobretamano(self):
        with self.assertRaises(UploadError) as cm:
            receive_upload(self._store(), filename="grande.csv",
                           size_bytes=DATA_UPLOAD_MAX_BYTES + 1,
                           content_iter=_chunks(b"x"), user_id="USR-1")
        self.assertIn("100 MB", str(cm.exception))

    def test_12_mensaje_sin_jerga(self):
        try:
            receive_upload(self._store(), filename="x.txt", size_bytes=10,
                           content_iter=_chunks(b"x"), user_id="USR-1")
        except UploadError as e:
            blob = str(e).lower()
            for word in ("parquet", "pipeline", "dataframe", "python", "json", "backend"):
                self.assertNotIn(word, blob)


class ValidationTest(unittest.TestCase):
    """Revisión de datos antes de procesar (unit)."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="zayvero7a_val_")
        self.orig_uploads = ds_store_mod.UPLOADS_DIR
        ds_store_mod.UPLOADS_DIR = os.path.join(self.tmp, "uploads")

    def tearDown(self):
        ds_store_mod.UPLOADS_DIR = self.orig_uploads
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _upload(self, data: bytes, filename="v.csv"):
        st = DatasetStore("CMP-VAL", base_dir=os.path.join(self.tmp, "datasets"))
        return receive_upload(st, filename=filename, size_bytes=len(data),
                              content_iter=_chunks(data), user_id="USR-1")

    def test_13_review_checks_basicos(self):
        ds = self._upload(_csv_bytes())
        rev = build_review(ds.upload["path"])
        self.assertTrue(rev["ok"])
        self.assertEqual(rev["total_rows"], 60)
        labels = [c["label"] for c in rev["checks"]]
        low = [l.lower() for l in labels]
        self.assertTrue(any("recibido" in l for l in low))
        self.assertTrue(any("columnas" in l for l in low))

    def test_14_review_columnas_faltantes(self):
        ds = self._upload(_csv_bytes())
        rev = build_review(ds.upload["path"])
        labels = [w["label"] for w in rev["warnings"]]
        self.assertIn("Columnas faltantes", labels)

    def test_15_review_negativos(self):
        ds = self._upload(_csv_bytes(negative=True))
        rev = build_review(ds.upload["path"])
        labels = [w["label"] for w in rev["warnings"]]
        self.assertIn("Valores negativos", labels)

    def test_16_review_duplicados(self):
        ds = self._upload(_csv_bytes(duplicates=True))
        rev = build_review(ds.upload["path"])
        labels = [w["label"] for w in rev["warnings"]]
        self.assertIn("Registros duplicados", labels)

    def test_17_review_cancelaciones(self):
        ds = self._upload(_csv_bytes(cancelled=True))
        rev = build_review(ds.upload["path"])
        labels = [w["label"] for w in rev["warnings"]]
        self.assertIn("Posibles cancelaciones", labels)

    def test_18_review_no_modifica_archivo(self):
        data = _csv_bytes()
        ds = self._upload(data)
        build_review(ds.upload["path"])
        with open(ds.upload["path"], "rb") as fh:
            self.assertEqual(fh.read(), data)

    def test_19_review_columnas_no_reconocidas_mensaje_veraz(self):
        # Las columnas no asignadas al esquema estándar se comunican sin
        # afirmar que "no se usarán": Tipo y Gasto_DOP se conservan como
        # orig_* y alimentan la detección de gastos.
        lines = ["Nota,Fecha,Tipo,Producto_o_concepto,Cantidad,"
                 "Precio_unitario_DOP,Gasto_DOP"]
        lines.append(",2026-09-01,Venta,Producto 1,2,387,0")
        lines.append(",2026-09-03,Gasto,Alquiler local,1,0,1500")
        data = ("\n".join(lines) + "\n").encode("utf-8")
        ds = self._upload(data)
        rev = build_review(ds.upload["path"])
        warns = [w for w in rev["warnings"]
                 if w["label"] == "Columnas no reconocidas"]
        self.assertEqual(len(warns), 1)
        detail = warns[0]["detail"]
        self.assertNotIn("no se usarán", detail.lower())
        self.assertIn("Tipo", detail)
        self.assertIn("Gasto_DOP", detail)
        self.assertIn("validaciones internas", detail)
        self.assertIn("Nota", detail)


class MappingTest(unittest.TestCase):
    """Mapeo de columnas (unit)."""

    COLS_ES = ["factura", "producto", "cantidad", "fecha_venta", "precio", "cliente"]

    def test_19_suggest_claro(self):
        sug = suggest_mapping(self.COLS_ES)
        by = {s["canonical"]: s for s in sug["suggestions"]}
        self.assertEqual(by["Quantity"]["source"], "cantidad")
        self.assertEqual(by["Quantity"]["status"], "suggested")
        self.assertEqual(by["UnitPrice"]["source"], "precio")
        self.assertEqual(by["InvoiceDate"]["source"], "fecha_venta")

    def test_20_suggest_ambiguo_requiere_revision(self):
        sug = suggest_mapping(self.COLS_ES)
        by = {s["canonical"]: s for s in sug["suggestions"]}
        # "producto" es ambiguo entre código y descripción → no se acepta solo.
        self.assertTrue(sug["needs_confirmation"])
        for canon in ("StockCode", "Description"):
            self.assertEqual(by[canon]["status"], "needs_review")

    def test_21_confirm_ok(self):
        mp = {"Invoice": "factura", "Quantity": "cantidad",
              "InvoiceDate": "fecha_venta", "UnitPrice": "precio",
              "StockCode": "producto"}
        res = confirm_mapping(self.COLS_ES, mp)
        self.assertTrue(res["ok"])
        self.assertEqual(res["mapping"]["Quantity"], "cantidad")

    def test_22_confirm_falta_requerido(self):
        res = confirm_mapping(self.COLS_ES, {"Invoice": "factura"})
        self.assertFalse(res["ok"])
        self.assertTrue(any("InvoiceDate" in e or "fecha" in e.lower()
                            for e in res["errors"]))

    def test_23_confirm_columna_inexistente(self):
        res = confirm_mapping(self.COLS_ES,
                              {"Invoice": "factura", "Quantity": "no_existe",
                               "InvoiceDate": "fecha_venta", "UnitPrice": "precio"})
        self.assertFalse(res["ok"])

    def test_24_confirm_sin_jerga(self):
        res = confirm_mapping(self.COLS_ES, {})
        blob = json.dumps(res["errors"]).lower()
        for word in ("parquet", "pipeline", "dataframe", "python", "json", "backend"):
            self.assertNotIn(word, blob)

    # --- FIX mapeo-selector: CSV dominicano de prueba (30 registros) ---
    COLS_DO = ["Tipo", "Producto_o_concepto", "Cantidad",
               "Precio_unitario_DOP", "Gasto_DOP", "Fecha"]

    def test_25_suggest_fecha_exacta(self):
        sug = suggest_mapping(self.COLS_DO)
        by = {s["canonical"]: s for s in sug["suggestions"]}
        self.assertEqual(by["InvoiceDate"]["source"], "Fecha")
        self.assertEqual(by["InvoiceDate"]["status"], "suggested")
        self.assertEqual(by["Quantity"]["source"], "Cantidad")

    def test_26_suggest_precio_dop(self):
        sug = suggest_mapping(self.COLS_DO)
        by = {s["canonical"]: s for s in sug["suggestions"]}
        self.assertEqual(by["UnitPrice"]["source"], "Precio_unitario_DOP")
        self.assertEqual(by["UnitPrice"]["status"], "suggested")

    def test_27_gasto_dop_no_es_precio(self):
        # Gasto_DOP es un concepto distinto: NO debe mapearse a UnitPrice
        # ni a ningún otro canónico; queda disponible sin asignar.
        sug = suggest_mapping(self.COLS_DO)
        for s in sug["suggestions"]:
            self.assertNotEqual(s["source"], "Gasto_DOP")

    def test_28_suggest_incluye_todas_las_columnas(self):
        # El selector debe mostrar TODAS las columnas originales, incluso
        # las desconocidas (Tipo, Producto_o_concepto, Gasto_DOP).
        sug = suggest_mapping(self.COLS_DO)
        self.assertEqual(sug["columns"], self.COLS_DO)
        by = {s["canonical"]: s for s in sug["suggestions"]}
        for canon in ("Invoice", "StockCode", "Description",
                      "CustomerID", "Country"):
            self.assertEqual(by[canon]["status"], "needs_review")
            self.assertIsNone(by[canon]["source"])

    def test_29_suggest_sin_duplicados(self):
        # Ninguna columna de origen se propone dos veces.
        sug = suggest_mapping(self.COLS_DO)
        srcs = [s["source"] for s in sug["suggestions"] if s["source"]]
        self.assertEqual(len(srcs), len(set(srcs)))

    def test_30_confirm_rechaza_duplicado(self):
        mp = {"Quantity": "Cantidad", "UnitPrice": "Cantidad",
              "InvoiceDate": "Fecha"}
        res = confirm_mapping(self.COLS_DO, mp)
        self.assertFalse(res["ok"])
        self.assertTrue(any("Cantidad" in e and "más de un campo" in e
                            for e in res["errors"]))

    def test_31_confirm_mapeo_do_ok(self):
        mp = {"Quantity": "Cantidad", "InvoiceDate": "Fecha",
              "UnitPrice": "Precio_unitario_DOP"}
        res = confirm_mapping(self.COLS_DO, mp)
        self.assertTrue(res["ok"])
        self.assertEqual(res["mapping"]["UnitPrice"], "Precio_unitario_DOP")
        # Las desconocidas simplemente quedan fuera del mapeo, sin error.
        self.assertNotIn("Gasto_DOP", res["mapping"].values())

    def test_32_online_retail_ii_intacto(self):
        # Compatibilidad: el esquema original sigue mapeando 1:1.
        cols = ["Invoice", "StockCode", "Description", "Quantity",
                "InvoiceDate", "UnitPrice", "CustomerID", "Country"]
        sug = suggest_mapping(cols)
        by = {s["canonical"]: s for s in sug["suggestions"]}
        for c in cols:
            self.assertEqual(by[c]["source"], c)
            self.assertEqual(by[c]["status"], "suggested")
        self.assertFalse(sug["needs_confirmation"])


class PreviewTest(unittest.TestCase):
    """Vista previa (unit)."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="zayvero7a_pv_")
        self.orig_uploads = ds_store_mod.UPLOADS_DIR
        ds_store_mod.UPLOADS_DIR = os.path.join(self.tmp, "uploads")

    def tearDown(self):
        ds_store_mod.UPLOADS_DIR = self.orig_uploads
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_25_preview_10_filas(self):
        st = DatasetStore("CMP-PV", base_dir=os.path.join(self.tmp, "datasets"))
        data = _csv_bytes(60)
        ds = receive_upload(st, filename="v.csv", size_bytes=len(data),
                            content_iter=_chunks(data), user_id="USR-1")
        pv = build_preview(ds.upload["path"])
        self.assertEqual(len(pv["rows"]), 10)
        self.assertEqual(pv["total_rows"], 60)
        self.assertTrue(pv["columns"])


class ProcessingTest(unittest.TestCase):
    """Procesamiento con el pipeline existente (unit)."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="zayvero7a_pr_")
        self.orig_uploads = ds_store_mod.UPLOADS_DIR
        ds_store_mod.UPLOADS_DIR = os.path.join(self.tmp, "uploads")
        self.tmp_tenant = tempfile.mkdtemp(prefix="zayvero7a_tenant_")
        self.tenant_store = TenantStore(self.tmp_tenant)
        self.comp = create_company(self.tenant_store, "Empresa Proc Test")
        self.cid = self.comp.company_id

    def tearDown(self):
        ds_store_mod.UPLOADS_DIR = self.orig_uploads
        shutil.rmtree(self.tmp, ignore_errors=True)
        shutil.rmtree(self.tmp_tenant, ignore_errors=True)
        _cleanup_company_data(self.cid)

    def _ready_dataset(self, nombre="Ventas Proc"):
        st = DatasetStore(self.cid)
        data = _csv_bytes(60)
        ds = receive_upload(st, filename="ventas.csv", size_bytes=len(data),
                            content_iter=_chunks(data), user_id="USR-1",
                            nombre=nombre)
        ds.review = build_review(ds.upload["path"])
        sug = suggest_mapping(ds.review["columns"])
        mp = {s["canonical"]: s["source"] for s in sug["suggestions"]
              if s["status"] == "suggested"}
        mp["InvoiceDate"] = "fecha_venta"
        mp["StockCode"] = "producto"
        res = confirm_mapping(ds.review["columns"], mp)
        self.assertTrue(res["ok"], res["errors"])
        from datasets.store import utcnow_iso
        ds.mapping = {"confirmed": True, "mapping": res["mapping"],
                      "confirmed_at": utcnow_iso(), "auto": False}
        st.save_dataset(ds)
        return st, ds

    def test_26_process_llega_a_ready(self):
        st, ds = self._ready_dataset()
        out = process_dataset(st, ds.dataset_id, tenant_store=self.tenant_store,
                              user_id="USR-1")
        self.assertEqual(out.status, "READY")
        self.assertTrue(out.is_active)  # primer READY se auto-activa
        self.assertIsNotNone(out.period_start)
        self.assertGreater(out.row_count, 0)
        for key in ("profile", "predictions", "business_context"):
            self.assertIn(key, out.outputs)
            self.assertTrue(os.path.isfile(out.outputs[key]),
                            f"falta output {key}")

    def test_27_process_etiqueta_nombre_cliente(self):
        st, ds = self._ready_dataset(nombre="Ventas Cliente XYZ")
        process_dataset(st, ds.dataset_id, tenant_store=self.tenant_store,
                        user_id="USR-1")
        prof_path = os.path.join(
            PROJECT_ROOT, "data", "profiles", self.cid,
            f"{ds.base_name}_profile.json")
        with open(prof_path, encoding="utf-8") as fh:
            prof = json.load(fh)
        label = prof["dataset_metadata"]["dataset_label"]
        self.assertEqual(label, "Ventas Cliente XYZ")
        self.assertNotIn("Demo", label)

    def test_28_process_error_mensaje_empresarial(self):
        st = DatasetStore(self.cid)
        data = b"esto,no,es,un,csv,valido\n1,2\n"
        ds = receive_upload(st, filename="malo.csv", size_bytes=len(data),
                            content_iter=_chunks(data), user_id="USR-1")
        from datasets.store import utcnow_iso
        ds.mapping = {"confirmed": True, "mapping": {}, "confirmed_at": utcnow_iso()}
        st.save_dataset(ds)
        out = process_dataset(st, ds.dataset_id, tenant_store=self.tenant_store,
                              user_id="USR-1")
        self.assertEqual(out.status, "ERROR")
        msg = (out.processing or {}).get("error", "")
        self.assertTrue(msg)
        for word in ("Traceback", "File \"", "parquet", "pipeline",
                     "dataframe", "python", "json", "backend", "/home/"):
            self.assertNotIn(word, msg)

    def test_29_audit_constantes_datasets(self):
        for const in ("DATASET_CREATED", "DATASET_UPLOAD_STARTED",
                      "DATASET_UPLOAD_COMPLETED", "DATASET_PROCESSING_STARTED",
                      "DATASET_PROCESSING_COMPLETED", "DATASET_PROCESSING_FAILED",
                      "DATASET_ACTIVATED"):
            self.assertTrue(hasattr(audit_mod, const), const)


# ===================== integración HTTP =====================

def _start_server(store):
    server_mod.STORE = store
    server_mod._DATA_CACHE.clear()
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), server_mod.WebappHandler)
    port = httpd.server_address[1]
    t = threading.Thread(target=httpd.serve_forever, daemon=True)
    t.start()
    return httpd, port


class Client:
    def __init__(self, port):
        self.port = port
        self.jar = http.cookiejar.CookieJar()

    def _req(self, method, path, body=None, raw_body=None, headers=None):
        opener = urllib.request.build_opener(
            urllib.request.HTTPCookieProcessor(self.jar))
        data = raw_body
        hdrs = dict(headers or {})
        if body is not None:
            data = json.dumps(body).encode("utf-8")
            hdrs["Content-Type"] = "application/json"
        req = urllib.request.Request(
            f"http://127.0.0.1:{self.port}{path}", data=data,
            headers=hdrs, method=method)
        try:
            with opener.open(req) as res:
                raw = res.read().decode("utf-8")
                try:
                    payload = json.loads(raw) if raw else {}
                except Exception:
                    payload = {"_raw": raw}
                return res.status, payload
        except urllib.error.HTTPError as e:
            raw = e.read().decode("utf-8")
            try:
                payload = json.loads(raw) if raw else {}
            except Exception:
                payload = {}
            return e.code, payload

    def get(self, path):
        return self._req("GET", path)

    def post(self, path, body=None):
        return self._req("POST", path, body=body)

    def login(self, email, password):
        return self.post("/api/login", {"email": email, "password": password})

    def upload(self, filename, content: bytes, nombre=""):
        boundary = "----ZAYVERO7A" + "x" * 16
        parts = []
        parts.append(f"--{boundary}\r\n"
                     f'Content-Disposition: form-data; name="nombre"\r\n\r\n'
                     f"{nombre}\r\n".encode())
        parts.append(f"--{boundary}\r\n"
                     f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
                     f"Content-Type: application/octet-stream\r\n\r\n".encode()
                     + content + b"\r\n")
        parts.append(f"--{boundary}--\r\n".encode())
        body = b"".join(parts)
        return self._req("POST", "/api/datasets/upload", raw_body=body,
                         headers={"Content-Type":
                                  f"multipart/form-data; boundary={boundary}"})


class DatasetsServerTest(unittest.TestCase):
    """Integración HTTP: workspace, upload, review, mapping, preview,
    process, activate, aislamiento, roles, auditoría."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="zayvero7a_srv_")
        store = TenantStore(cls.tmp)
        store.ensure_demo_company()
        cls.comp_a = create_company(store, "Empresa A 7A")
        cls.comp_b = create_company(store, "Empresa B 7A")
        for email, cid, role in [
            ("ownera@7a.test", cls.comp_a.company_id, "owner"),
            ("admina@7a.test", cls.comp_a.company_id, "admin"),
            ("viewera@7a.test", cls.comp_a.company_id, "viewer"),
            ("ownerb@7a.test", cls.comp_b.company_id, "owner"),
        ]:
            create_user(store, company_id=cid, email=email,
                        name=email.split("@")[0], password=PASS, role_id=role)
        create_user(store, company_id="demo-retail", email="owner@demo.test",
                    name="Demo", password=PASS, role_id="owner")
        cls.store = store
        cls.httpd, cls.port = _start_server(store)
        cls._companies = [cls.comp_a.company_id, cls.comp_b.company_id]

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        for cid in cls._companies:
            _cleanup_company_data(cid)
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def setUp(self):
        self.c = Client(self.port)

    def _login_a(self, email="ownera@7a.test"):
        code, _ = self.c.login(email, PASS)
        self.assertEqual(code, 200)

    # ---- workspace -------------------------------------------------
    def test_30_workspace_estructura(self):
        self._login_a()
        code, ws = self.c.get("/api/workspace")
        self.assertEqual(code, 200)
        for key in ("company", "user", "dataset_state", "dataset_state_message",
                    "datasets", "can_upload"):
            self.assertIn(key, ws, key)
        self.assertEqual(ws["company"]["name"], "Empresa A 7A")
        self.assertEqual(ws["user"]["role"], "owner")

    def test_31_workspace_no_data(self):
        self._login_a()
        _, ws = self.c.get("/api/workspace")
        self.assertEqual(ws["dataset_state"], "NO_DATA")
        self.assertIn("datos", ws["dataset_state_message"].lower())
        self.assertEqual(ws["datasets"], [])
        self.assertTrue(ws["can_upload"])

    def test_32_workspace_viewer_no_puede_subir(self):
        self.c.login("viewera@7a.test", PASS)
        _, ws = self.c.get("/api/workspace")
        self.assertFalse(ws["can_upload"])

    def test_33_workspace_demo_bandera(self):
        self.c.login("owner@demo.test", PASS)
        code, ws = self.c.get("/api/workspace")
        self.assertEqual(code, 200)
        self.assertTrue(ws["company"]["is_demo"])

    # ---- upload ----------------------------------------------------
    def test_34_upload_owner_ok(self):
        self._login_a()
        code, res = self.c.upload("ventas.csv", _csv_bytes(60), "Ventas 7A")
        self.assertEqual(code, 200, res)
        ds = res["dataset"]
        self.assertEqual(ds["status"], "DATA_UPLOADED")
        self.assertEqual(ds["company_id"], self.comp_a.company_id)
        self.assertEqual(ds["upload"]["format"], "csv")
        self.assertEqual(ds["upload"]["filename"], "ventas.csv")
        self.__class__._ds_a = ds["dataset_id"]

    def test_35_upload_viewer_403(self):
        self.c.login("viewera@7a.test", PASS)
        code, _ = self.c.upload("v.csv", _csv_bytes(10))
        self.assertEqual(code, 403)

    def test_36_upload_admin_ok(self):
        self.c.login("admina@7a.test", PASS)
        code, res = self.c.upload("admin.csv", _csv_bytes(10), "Admin")
        self.assertEqual(code, 200, res)

    def test_37_upload_formato_invalido_400(self):
        self._login_a()
        code, res = self.c.upload("reporte.pdf", b"%PDF-1.4")
        self.assertEqual(code, 400)
        self.assertIn("CSV", res.get("error", ""))

    def test_38_upload_empresa_b_aislada(self):
        self.c.login("ownerb@7a.test", PASS)
        code, res = self.c.upload("b.csv", _csv_bytes(20), "Ventas B")
        self.assertEqual(code, 200, res)
        self.assertEqual(res["dataset"]["company_id"], self.comp_b.company_id)
        _, ws = self.c.get("/api/workspace")
        self.assertEqual(len(ws["datasets"]), 1)

    # ---- review / mapping / preview --------------------------------
    def _ensure_ds_a(self):
        ds_id = getattr(self.__class__, "_ds_a", None)
        if ds_id is None:
            self._login_a()
            _, res = self.c.upload("ventas.csv", _csv_bytes(60), "Ventas 7A")
            ds_id = res["dataset"]["dataset_id"]
            self.__class__._ds_a = ds_id
        return ds_id

    def test_39_review_endpoint(self):
        self._login_a()
        ds_id = self._ensure_ds_a()
        code, res = self.c.get(f"/api/datasets/{ds_id}/review")
        self.assertEqual(code, 200, res)
        rev = res["review"]
        self.assertTrue(rev["ok"])
        self.assertEqual(rev["total_rows"], 60)
        self.assertTrue(any("Columnas faltantes" in w["label"]
                            for w in rev["warnings"]))

    def test_40_mapping_sugerencias(self):
        self._login_a()
        ds_id = self._ensure_ds_a()
        code, res = self.c.get(f"/api/datasets/{ds_id}/mapping")
        self.assertEqual(code, 200, res)
        sug = res["suggestion"]
        by = {s["canonical"]: s for s in sug["suggestions"]}
        self.assertEqual(by["Quantity"]["source"], "cantidad")
        self.assertTrue(sug["needs_confirmation"])

    def test_41_mapping_confirm_incompleto_400(self):
        self._login_a()
        ds_id = self._ensure_ds_a()
        code, res = self.c.post(f"/api/datasets/{ds_id}/mapping",
                                {"mapping": {"Invoice": "factura"}})
        self.assertEqual(code, 400)
        self.assertTrue(res.get("errors"))

    def test_42_mapping_confirm_ok(self):
        self._login_a()
        ds_id = self._ensure_ds_a()
        mp = {"Invoice": "factura", "StockCode": "producto",
              "Quantity": "cantidad", "InvoiceDate": "fecha_venta",
              "UnitPrice": "precio", "CustomerID": "cliente"}
        code, res = self.c.post(f"/api/datasets/{ds_id}/mapping",
                                {"mapping": mp})
        self.assertEqual(code, 200, res)
        self.assertTrue(res["mapping"]["confirmed"])

    def test_43_preview_endpoint(self):
        self._login_a()
        ds_id = self._ensure_ds_a()
        code, res = self.c.get(f"/api/datasets/{ds_id}/preview")
        self.assertEqual(code, 200, res)
        pv = res["preview"]
        self.assertEqual(len(pv["rows"]), 10)
        self.assertEqual(pv["total_rows"], 60)

    # ---- aislamiento ----------------------------------------------
    def test_44_b_no_ve_dataset_de_a(self):
        self.c.login("ownerb@7a.test", PASS)
        ds_id = self._ensure_ds_a()
        code, _ = self.c.get(f"/api/datasets/{ds_id}")
        self.assertEqual(code, 403)
        code, _ = self.c.get(f"/api/datasets/{ds_id}/review")
        self.assertEqual(code, 403)
        code, _ = self.c.post(f"/api/datasets/{ds_id}/mapping", {"mapping": {}})
        self.assertEqual(code, 403)

    def test_45_b_no_ve_lista_de_a(self):
        self.c.login("ownerb@7a.test", PASS)
        _, res = self.c.get("/api/datasets")
        ids = [d["dataset_id"] for d in res["datasets"]]
        self.assertNotIn(self._ensure_ds_a(), ids)

    # ---- process / activate ----------------------------------------
    def test_46_process_202_y_ready(self):
        self._login_a()
        ds_id = self._ensure_ds_a()
        # Confirmar mapeo por si el orden de tests varió.
        mp = {"Invoice": "factura", "StockCode": "producto",
              "Quantity": "cantidad", "InvoiceDate": "fecha_venta",
              "UnitPrice": "precio", "CustomerID": "cliente"}
        self.c.post(f"/api/datasets/{ds_id}/mapping", {"mapping": mp})
        code, res = self.c.post(f"/api/datasets/{ds_id}/process", {})
        self.assertIn(code, (200, 202), res)
        # Esperar READY (procesamiento en segundo plano).
        deadline = time.time() + 120
        status = None
        while time.time() < deadline:
            _, det = self.c.get(f"/api/datasets/{ds_id}")
            status = det["dataset"]["status"]
            if status in ("READY", "ERROR"):
                break
            time.sleep(2)
        self.assertEqual(status, "READY", det["dataset"].get("processing"))

    def test_47_dashboard_ready_con_datos_de_a(self):
        self._login_a()
        code, s = self.c.get("/api/summary")
        self.assertEqual(code, 200)
        self.assertTrue(s.get("ready"))
        self.assertEqual(s.get("dataset_state"), "READY")
        blob = json.dumps(s).lower()
        self.assertNotIn("demo-retail", blob)
        self.assertNotIn("uci online retail", blob)

    def test_48_workspace_ready_despues_de_proceso(self):
        self._login_a()
        _, ws = self.c.get("/api/workspace")
        self.assertEqual(ws["dataset_state"], "READY")
        ad = ws["active_dataset"]
        self.assertIsNotNone(ad)
        self.assertGreater(ad["row_count"], 0)
        self.assertTrue(ad["period_start"])
        self.assertTrue(ad["period_end"])

    def test_49_reprocess_ok(self):
        self._login_a()
        ds_id = self._ensure_ds_a()
        code, res = self.c.post(f"/api/datasets/{ds_id}/process", {})
        self.assertIn(code, (200, 202), res)
        deadline = time.time() + 120
        status = None
        while time.time() < deadline:
            _, det = self.c.get(f"/api/datasets/{ds_id}")
            status = det["dataset"]["status"]
            if status in ("READY", "ERROR"):
                break
            time.sleep(2)
        self.assertEqual(status, "READY")

    def test_50_activate_segundo_dataset(self):
        self._login_a()
        _, res = self.c.upload("segundo.csv", _csv_bytes(40), "Segundo")
        ds2 = res["dataset"]["dataset_id"]
        mp = {"Invoice": "factura", "StockCode": "producto",
              "Quantity": "cantidad", "InvoiceDate": "fecha_venta",
              "UnitPrice": "precio", "CustomerID": "cliente"}
        self.c.post(f"/api/datasets/{ds2}/mapping", {"mapping": mp})
        code, _ = self.c.post(f"/api/datasets/{ds2}/activate", {})
        self.assertEqual(code, 400)  # no está READY → no se puede activar
        # Activar el dataset READY (idempotente).
        ds1 = self._ensure_ds_a()
        code, res = self.c.post(f"/api/datasets/{ds1}/activate", {})
        self.assertEqual(code, 200)
        self.assertTrue(res["dataset"]["is_active"])

    # ---- advisor / demo / auditoría ---------------------------------
    def test_51_advisor_solo_contexto_de_a(self):
        self._login_a()
        code, res = self.c.post("/api/advisor/ask",
                                {"question": "¿Cómo va mi empresa?"})
        self.assertEqual(code, 200)
        blob = json.dumps(res).lower()
        self.assertNotIn("demo-retail", blob)
        self.assertNotIn("uci online retail", blob)

    def test_52_advisor_b_sin_datos_no_inventa(self):
        self.c.login("ownerb@7a.test", PASS)
        code, res = self.c.post("/api/advisor/ask",
                                {"question": "¿Cuánto vendí en enero?"})
        self.assertEqual(code, 200)
        self.assertTrue(res.get("fallback"))
        self.assertIn("todavía no tengo datos", res.get("answer", "").lower())

    def test_53_demo_aislada_de_a(self):
        # A no ve los datasets del demo.
        self._login_a()
        _, res = self.c.get("/api/datasets")
        self.assertFalse(any(d.get("is_demo") for d in res["datasets"]))

    def test_54_auditoria_eventos_datasets(self):
        self._login_a()
        code, res = self.c.get("/api/audit")
        self.assertEqual(code, 200)
        actions = [e["action"] for e in res["events"]]
        for expected in ("DATASET_CREATED", "DATASET_UPLOAD_STARTED",
                         "DATASET_UPLOAD_COMPLETED"):
            self.assertIn(expected, actions, expected)
        # Sin secretos en la auditoría.
        blob = json.dumps(res).lower()
        for word in ("password", "contraseña", "token", "secret", "pbkdf2"):
            self.assertNotIn(word, blob)

    def test_55_estados_dashboard_no_data_b(self):
        self.c.login("ownerb@7a.test", PASS)
        code, s = self.c.get("/api/summary")
        self.assertEqual(code, 200)
        # B tiene un dataset cargado pero no procesado → no READY.
        self.assertFalse(s.get("ready"))
        self.assertIn(s.get("dataset_state"),
                      ("NO_DATA", "DATA_UPLOADED", "PROCESSING"))

    def test_56_lenguaje_empresarial_estado(self):
        self._login_a()
        _, ws = self.c.get("/api/workspace")
        blob = json.dumps(ws).lower()
        for word in ("parquet", "pipeline", "dataframe", "python", "backend"):
            self.assertNotIn(word, blob)


if __name__ == "__main__":
    unittest.main()
